"""Read-only structural extraction; NOT ERC/DRC or a connectivity verifier.
Run from repository root with Python + sexpdata, pypdf, openpyxl installed.
Only writes ROXX_Flight_Controller/audit. Original inputs are git-tracked files.
"""
from pathlib import Path
import csv, hashlib, json, subprocess
import sexpdata as sx
from pypdf import PdfReader
from openpyxl import load_workbook
ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / 'ROXX_Flight_Controller/audit'
OUT.mkdir(exist_ok=True)
def nodes(obj, key):
    return [x for x in obj if isinstance(x, list) and x and str(x[0]) == key]
def table(name, header, rows):
    with (OUT / name).open('w', newline='') as f:
        w = csv.writer(f); w.writerow(header); w.writerows(rows)
paths = [Path(p) for p in subprocess.check_output(['git','ls-tree','-rz','--name-only','HEAD'],cwd=ROOT,text=True).split('\0') if p and not p.startswith('ROXX_Flight_Controller/')]
table('source_manifest.csv', ['path','bytes','sha256'], [(str(p),(ROOT/p).stat().st_size,hashlib.sha256((ROOT/p).read_bytes()).hexdigest()) for p in paths])
sch = sx.load((ROOT/'NxtFC-DUAL/NxtPX4.kicad_sch').open())
pcb = sx.load((ROOT/'NxtFC-DUAL/NxtPX4.kicad_pcb').open())
symbols = {}
for obj in nodes(sch,'symbol'):
    props = {p[1]:p[2] for p in nodes(obj,'property')}
    ref = props['Reference']
    symbols[ref] = [ref, props.get('Value',''),props.get('Footprint',''),props.get('Datasheet',''), nodes(obj,'lib_id')[0][1]]
table('reference_symbols.csv',['reference','value','footprint','datasheet_field_unverified','library_id'],sorted(symbols.values()))
pads=[]; footprints={}; models=[]
for obj in nodes(pcb,'footprint'):
    texts={str(p[1]):p[2] for p in nodes(obj,'fp_text')}
    ref=texts.get('reference','')
    footprints[ref]=[ref,texts.get('value',''),obj[1],str(nodes(obj,'at'))]
    for pad in nodes(obj,'pad'):
        nets=nodes(pad,'net')
        pads.append([ref,pad[1],nets[0][2] if nets else '',str(nodes(pad,'size')),str(nodes(pad,'drill'))])
    for m in nodes(obj,'model'):
        path=m[1]; local=path.replace('${KIPRJMOD}',str(ROOT/'NxtFC-DUAL'))
        models.append([ref,path,'external-library-unchecked' if '${' in local else ('present' if Path(local).exists() else 'missing')])
table('reference_pcb_footprints.csv',['reference','value','library_id','position'],sorted(footprints.values()))
table('reference_pcb_pad_nets.csv',['reference','pad','net','size','drill'],pads)
table('reference_model_paths.csv',['reference','model','resolution_status'],models)
summary={'scope':'Structural extraction only; not ERC/DRC', 'source_commit':subprocess.check_output(['git','rev-parse','HEAD'],cwd=ROOT,text=True).strip(), 'source_files':len(paths),'schematic_unique_references_including_power':len(symbols),'pcb_footprints':len(footprints),'schematic_only_non_power':sorted(r for r in symbols if r not in footprints and not r.startswith('#')), 'pcb_only':sorted(set(footprints)-set(symbols)), 'value_differences':[[r,symbols[r][1],v[1]] for r,v in footprints.items() if r in symbols and symbols[r][1]!=v[1]]}
(OUT/'structural_summary.json').write_text(json.dumps(summary,indent=2)+'\n')
for p in paths:
    if p.suffix.lower()=='.pdf':
        reader=PdfReader(ROOT/p)
        (OUT/(p.stem+'_text.txt')).write_text('SOURCE: '+str(p)+'\nPAGES: '+str(len(reader.pages))+'\n'+'\n'.join(page.extract_text() for page in reader.pages))
w=load_workbook(ROOT/'NXTVersion2.0/STM32H743VIH_pinout_v2.0.xlsx',data_only=True)
with (OUT/'reference_pin_workbook.txt').open('w') as f:
    for sh in w:
        f.write('\nSHEET: '+sh.title+'\n')
        for row in sh.iter_rows():
            vals=[str(c.value) if c.value is not None else '' for c in row]
            if any(vals): f.write('\t'.join(vals)+'\n')
print(json.dumps(summary,indent=2))
